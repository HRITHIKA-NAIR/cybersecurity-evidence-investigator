from io import BytesIO

from openpyxl import load_workbook

MAX_SPREADSHEET_CELLS = 20_000


def parse_excel(data: bytes) -> str:
    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=False,
    )
    parts = []
    cells_seen = 0

    try:
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(
                values_only=True
            ):
                values = [
                    str(value)
                    for value in row
                    if value is not None
                ]
                cells_seen += len(row)

                if values:
                    parts.append(
                        " | ".join(values)
                    )

                if (
                    cells_seen
                    >= MAX_SPREADSHEET_CELLS
                ):
                    parts.append(
                        "[Spreadsheet extraction "
                        "limit reached]"
                    )
                    return "\n".join(parts)
    finally:
        workbook.close()

    return "\n".join(parts)
