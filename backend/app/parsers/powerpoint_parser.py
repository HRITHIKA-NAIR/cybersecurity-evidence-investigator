from io import BytesIO

from pptx import Presentation


def parse_powerpoint(data: bytes) -> str:
    presentation = Presentation(BytesIO(data))
    parts = []

    for number, slide in enumerate(
        presentation.slides,
        start=1,
    ):
        text = [
            shape.text.strip()
            for shape in slide.shapes
            if (
                hasattr(shape, "text")
                and shape.text.strip()
            )
        ]

        if text:
            parts.append(
                f"Slide {number}\n"
                + "\n".join(text)
            )

    return "\n\n".join(parts)
