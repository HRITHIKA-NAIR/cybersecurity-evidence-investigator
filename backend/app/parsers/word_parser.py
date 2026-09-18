from io import BytesIO

from docx import Document


def parse_word(data: bytes) -> str:
    document = Document(BytesIO(data))
    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            values = [
                cell.text.strip()
                for cell in row.cells
            ]

            if any(values):
                parts.append(" | ".join(values))

    return "\n".join(parts)
