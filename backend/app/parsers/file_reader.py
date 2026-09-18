from __future__ import annotations

import hashlib
import zipfile
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path

MAX_UPLOAD_BYTES = 10 * 1024 * 1024
MAX_EXTRACTED_CHARS = 50_000
MAX_ARCHIVE_ITEMS = 500
MAX_ARCHIVE_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_COMPRESSION_RATIO = 100
MAX_SPREADSHEET_CELLS = 20_000

TEXT_EXTENSIONS = {
    ".txt",
    ".md",
    ".csv",
    ".json",
    ".html",
    ".htm",
    ".svg",
}
SUPPORTED_EXTENSIONS = TEXT_EXTENSIONS | {
    ".eml",
    ".pdf",
    ".docx",
    ".pptx",
    ".xlsx",
    ".zip",
}
BINARY_EXTENSIONS = {".pdf", ".docx", ".pptx", ".xlsx", ".zip"}


class FileReaderError(ValueError):
    def __init__(self, message: str, status_code: int = 400):
        super().__init__(message)
        self.status_code = status_code


def _truncate(text: str) -> tuple[str, bool]:
    text = text.strip()
    if len(text) <= MAX_EXTRACTED_CHARS:
        return text, False
    return text[:MAX_EXTRACTED_CHARS], True


def _decode_text(data: bytes) -> str:
    return data.decode("utf-8-sig", errors="replace")


def _inspect_zip(data: bytes) -> list[zipfile.ZipInfo]:
    with zipfile.ZipFile(BytesIO(data)) as archive:
        infos = archive.infolist()

    if len(infos) > MAX_ARCHIVE_ITEMS:
        raise FileReaderError(
            f"Archive contains too many items (maximum {MAX_ARCHIVE_ITEMS}).",
            413,
        )

    total_uncompressed = sum(info.file_size for info in infos)
    if total_uncompressed > MAX_ARCHIVE_UNCOMPRESSED_BYTES:
        raise FileReaderError(
            "Archive expands beyond the safe processing limit.",
            413,
        )

    for info in infos:
        ratio = info.file_size / max(info.compress_size, 1)
        if ratio > MAX_COMPRESSION_RATIO and info.file_size > 1024 * 1024:
            raise FileReaderError(
                "Archive has an unsafe compression ratio.",
                413,
            )

    return infos


def _detect_type(data: bytes) -> str:
    if data.startswith(b"%PDF-"):
        return ".pdf"

    if zipfile.is_zipfile(BytesIO(data)):
        infos = _inspect_zip(data)
        names = [info.filename for info in infos]

        if any(name.startswith("word/") for name in names):
            return ".docx"
        if any(name.startswith("ppt/") for name in names):
            return ".pptx"
        if any(name.startswith("xl/") for name in names):
            return ".xlsx"
        return ".zip"

    return "text"


def _read_eml(data: bytes) -> str:
    message = BytesParser(policy=policy.default).parsebytes(data)
    parts = [
        f"From: {message.get('From', '')}",
        f"To: {message.get('To', '')}",
        f"Subject: {message.get('Subject', '')}",
        f"Date: {message.get('Date', '')}",
    ]

    body = message.get_body(preferencelist=("plain", "html"))
    if body:
        try:
            parts.append(body.get_content())
        except (KeyError, LookupError, UnicodeDecodeError):
            pass

    return "\n".join(part for part in parts if part.strip())


def _read_pdf(data: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(data))
    return "\n\n".join(
        text
        for page in reader.pages
        if (text := page.extract_text())
    )


def _read_docx(data: bytes) -> str:
    from docx import Document

    document = Document(BytesIO(data))
    parts = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    for table in document.tables:
        for row in table.rows:
            values = [cell.text.strip() for cell in row.cells]
            if any(values):
                parts.append(" | ".join(values))

    return "\n".join(parts)


def _read_pptx(data: bytes) -> str:
    from pptx import Presentation

    presentation = Presentation(BytesIO(data))
    parts = []

    for number, slide in enumerate(presentation.slides, start=1):
        text = [
            shape.text.strip()
            for shape in slide.shapes
            if hasattr(shape, "text") and shape.text.strip()
        ]
        if text:
            parts.append(f"Slide {number}\n" + "\n".join(text))

    return "\n\n".join(parts)


def _read_xlsx(data: bytes) -> str:
    from openpyxl import load_workbook

    workbook = load_workbook(
        BytesIO(data),
        read_only=True,
        data_only=False,
    )
    parts = []
    cells_seen = 0

    try:
        for sheet in workbook.worksheets:
            parts.append(f"Sheet: {sheet.title}")

            for row in sheet.iter_rows(values_only=True):
                values = [str(value) for value in row if value is not None]
                cells_seen += len(row)

                if values:
                    parts.append(" | ".join(values))

                if cells_seen >= MAX_SPREADSHEET_CELLS:
                    parts.append("[Spreadsheet extraction limit reached]")
                    return "\n".join(parts)
    finally:
        workbook.close()

    return "\n".join(parts)


def _read_zip(data: bytes) -> str:
    infos = _inspect_zip(data)
    return "Archive contents:\n" + "\n".join(
        info.filename for info in infos
    )


def read_uploaded_file(
    filename: str,
    content_type: str | None,
    data: bytes,
) -> dict:
    safe_name = Path(filename or "upload").name
    extension = Path(safe_name).suffix.lower()

    if not data:
        raise FileReaderError("The uploaded file is empty.")

    if len(data) > MAX_UPLOAD_BYTES:
        raise FileReaderError(
            "File is too large. Maximum upload size is 10 MB.",
            413,
        )

    if extension not in SUPPORTED_EXTENSIONS:
        raise FileReaderError(
            f"Unsupported file type: {extension or 'no extension'}.",
            415,
        )

    detected_type = _detect_type(data)

    if extension in BINARY_EXTENSIONS and detected_type != extension:
        raise FileReaderError(
            "The file content does not match its extension.",
            415,
        )

    if extension in TEXT_EXTENSIONS | {".eml"} and detected_type != "text":
        raise FileReaderError(
            "The file content does not match its extension.",
            415,
        )

    readers = {
        ".eml": _read_eml,
        ".pdf": _read_pdf,
        ".docx": _read_docx,
        ".pptx": _read_pptx,
        ".xlsx": _read_xlsx,
        ".zip": _read_zip,
    }

    try:
        if extension in TEXT_EXTENSIONS:
            text = _decode_text(data)
            parser = "text"
        else:
            text = readers[extension](data)
            parser = extension.lstrip(".")
    except FileReaderError:
        raise
    except Exception as exc:
        raise FileReaderError(
            f"Could not read {extension} file.",
            422,
        ) from exc

    content, truncated = _truncate(text)

    return {
        "content": content,
        "file_info": {
            "filename": safe_name,
            "extension": extension,
            "declared_mime": content_type,
            "detected_type": detected_type,
            "size_bytes": len(data),
            "sha256": hashlib.sha256(data).hexdigest(),
            "parser": parser,
            "characters_extracted": len(content),
            "truncated": truncated,
        },
    }
