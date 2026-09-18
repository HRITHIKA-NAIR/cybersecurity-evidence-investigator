from io import BytesIO
from zipfile import ZipFile

import pytest

from app.parsers.file_reader import (
    FileReaderError,
    MAX_UPLOAD_BYTES,
    read_uploaded_file,
)


def test_reads_text():
    result = read_uploaded_file(
        "note.txt",
        "text/plain",
        b"hello https://example.com",
    )
    assert "https://example.com" in result["content"]
    assert result["file_info"]["extension"] == ".txt"


def test_rejects_empty_file():
    with pytest.raises(FileReaderError) as exc:
        read_uploaded_file("empty.txt", "text/plain", b"")

    assert exc.value.status_code == 400


def test_rejects_unsupported_extension():
    with pytest.raises(FileReaderError) as exc:
        read_uploaded_file("sample.exe", None, b"MZtest")

    assert exc.value.status_code == 415


def test_rejects_oversized_file():
    with pytest.raises(FileReaderError) as exc:
        read_uploaded_file(
            "large.txt",
            "text/plain",
            b"a" * (MAX_UPLOAD_BYTES + 1),
        )

    assert exc.value.status_code == 413


def test_rejects_binary_extension_mismatch():
    with pytest.raises(FileReaderError) as exc:
        read_uploaded_file(
            "fake.pdf",
            "application/pdf",
            b"not actually a pdf",
        )

    assert exc.value.status_code == 415


def test_reads_eml_body():
    data = (
        b"From: Sender <sender@example.com>\r\n"
        b"To: user@example.org\r\n"
        b"Subject: Verify account\r\n"
        b"Date: Tue, 15 Sep 2026 10:00:00 +0000\r\n"
        b"\r\n"
        b"Visit https://example.com/verify\r\n"
    )
    result = read_uploaded_file("message.eml", "message/rfc822", data)

    assert "Verify account" in result["content"]
    assert "https://example.com/verify" in result["content"]


def test_reads_docx():
    from docx import Document

    buffer = BytesIO()
    document = Document()
    document.add_paragraph("Document evidence")
    document.save(buffer)

    result = read_uploaded_file(
        "report.docx",
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        buffer.getvalue(),
    )

    assert "Document evidence" in result["content"]
    assert result["file_info"]["detected_type"] == ".docx"


def test_reads_pptx():
    from pptx import Presentation

    buffer = BytesIO()
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    textbox = slide.shapes.add_textbox(0, 0, 100, 100)
    textbox.text = "Presentation evidence"
    presentation.save(buffer)

    result = read_uploaded_file(
        "slides.pptx",
        "application/vnd.openxmlformats-officedocument.presentationml.presentation",
        buffer.getvalue(),
    )

    assert "Presentation evidence" in result["content"]


def test_reads_xlsx():
    from openpyxl import Workbook

    buffer = BytesIO()
    workbook = Workbook()
    workbook.active["A1"] = "Spreadsheet evidence"
    workbook.save(buffer)

    result = read_uploaded_file(
        "sheet.xlsx",
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        buffer.getvalue(),
    )

    assert "Spreadsheet evidence" in result["content"]


def test_lists_zip_without_extracting():
    buffer = BytesIO()
    with ZipFile(buffer, "w") as archive:
        archive.writestr("invoice.txt", "content")
        archive.writestr("nested/file.txt", "content")

    result = read_uploaded_file(
        "archive.zip",
        "application/zip",
        buffer.getvalue(),
    )

    assert "invoice.txt" in result["content"]
    assert "nested/file.txt" in result["content"]
