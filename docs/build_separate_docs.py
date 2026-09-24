"""Build fourteen independent Word documents from the maintained Markdown sources."""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / 'docs' / 'v2.3_separate'
ITEMS = [
 ('PRD','PRD.md','Product Requirements Document'),
 ('DESIGN_DOC','DESIGN.md','Interface Design Document'),
 ('SRS','SRS.md','Software Requirements Specification'),
 ('TDD','TDD.md','Technical and Test Design Document'),
 ('TECH_STACK','TECH_STACK.md','Technology Stack Document'),
 ('SDD','SDD.md','Software Design Document'),
 ('USER_GUIDE','USER_GUIDE.md','User Guide'),
 ('DEVELOPER_GUIDE','DEVELOPER_GUIDE.md','Developer Guide and Launch Steps'),
 ('WIREFRAMES','WIREFRAMES.md','Wireframes and Interface States'),
 ('PROJECT_HISTORY','PROJECT_HISTORY.md','Project History and Completion Record'),
 ('DESIGN_SYSTEM','DESIGN_SYSTEM.md','Unified Design System'),
 ('FEATURE_LIST','FEATURE_LIST.md','Feature List'),
 ('BRAND_GUIDELINES','BRAND_GUIDELINES.md','Brand and Design Guidelines'),
 ('APP_CHARTER','APP_CHARTER.md','App Charter'),
]


def inline(p,text):
    for part in re.split(r'(\*\*[^*]+\*\*|`[^`]+`|\[[^\]]+\]\([^)]+\))',text):
        if not part: continue
        link=re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)',part)
        if link:
            h=OxmlElement('w:hyperlink');h.set(qn('r:id'),p.part.relate_to(link[2],RT.HYPERLINK,is_external=True))
            r=OxmlElement('w:r');pr=OxmlElement('w:rPr');c=OxmlElement('w:color');c.set(qn('w:val'),'164F73');pr.append(c);r.append(pr)
            t=OxmlElement('w:t');t.text=link[1];r.append(t);h.append(r);p._p.append(h)
        else:
            bold=part.startswith('**') and part.endswith('**')
            code=part.startswith('`') and part.endswith('`')
            run=p.add_run(part[2:-2] if bold else part[1:-1] if code else part)
            run.bold=bold
            if code: run.font.name='Consolas';run.font.size=Pt(9.5)


def table(doc,rows):
    if not rows:return
    n=len(rows[0]);t=doc.add_table(rows=0,cols=n);t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    # Short identifiers and route names need less space than narrative fields.
    if n==4 and rows[0][0]=='Token': widths=[1.65,1.15,1.15,2.9]
    elif n==4 and rows[0][0]=='Style': widths=[1.4,1.55,1.0,2.9]
    elif n==4: widths=[.9,1.45,2.8,1.7]
    elif n==3: widths=[1.7,2.15,3.0]
    elif n==2: widths=[2.15,4.7]
    else: widths=[6.85/n]*n
    for col,w in zip(t.columns,widths):col.width=Inches(w)
    pr=t._tbl.tblPr;bd=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'4');el.set(qn('w:color'),'D9D9D9');bd.append(el)
    pr.append(bd)
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        rowpr=t.rows[-1]._tr.get_or_add_trPr();cant=OxmlElement('w:cantSplit');rowpr.append(cant)
        if i==0:rowpr.append(OxmlElement('w:tblHeader'))
        for j,cell in enumerate(cells):
            cell.width=Inches(widths[j]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p=cell.paragraphs[0];p.paragraph_format.space_before=Pt(4);p.paragraph_format.space_after=Pt(5);p.paragraph_format.line_spacing=1.06
            inline(p,row[j] if j<len(row) else '')
            for r in p.runs:r.font.size=Pt(9.5);r.bold=i==0;r.font.color.rgb=RGBColor.from_string('FFFFFF' if i==0 else '000000')
            cellpr=cell._tc.get_or_add_tcPr();shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'284A60' if i==0 else 'F1F5F7' if i%2==0 else 'FFFFFF');cellpr.append(shade)
            mar=OxmlElement('w:tcMar')
            for edge in ('top','left','bottom','right'):
                el=OxmlElement('w:'+edge);el.set(qn('w:w'),'85');el.set(qn('w:type'),'dxa');mar.append(el)
            cellpr.append(mar)
    doc.add_paragraph().paragraph_format.space_after=Pt(1)


def markdown(doc,text):
    lines=text.splitlines();i=0
    while i<len(lines):
        line=lines[i];i+=1
        if not line.strip():continue
        if line.startswith('```'):
            code=[]
            while i<len(lines) and not lines[i].startswith('```'):code.append(lines[i]);i+=1
            i+=1
            for value in code:
                p=doc.add_paragraph(style='No Spacing');p.paragraph_format.space_after=Pt(1)
                r=p.add_run(value);r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
                fonts=r._element.get_or_add_rPr().rFonts
                for attr in ('asciiTheme','hAnsiTheme','eastAsiaTheme','cstheme'):
                    fonts.attrib.pop(qn('w:'+attr),None)
            doc.add_paragraph().paragraph_format.space_after=Pt(1)
            continue
        picture=re.fullmatch(r'!\[([^\]]*)\]\(([^)]+)\)',line)
        if picture:
            p=doc.add_paragraph();p.paragraph_format.keep_with_next=True
            width=1.0 if 'brand-shield' in picture[2] else 6.85
            shape=p.add_run().add_picture(str(ROOT/'docs'/picture[2]),width=Inches(width));shape._inline.docPr.set('descr',picture[1])
            cap=doc.add_paragraph(picture[1],style='Caption');cap.paragraph_format.space_after=Pt(10)
            continue
        if line.startswith('|'):
            rows=[]
            while True:
                row=[v.strip() for v in line.strip().strip('|').split('|')]
                if not all(re.fullmatch(r'[-: ]+',v) for v in row):rows.append(row)
                if i>=len(lines) or not lines[i].startswith('|'):break
                line=lines[i];i+=1
            table(doc,rows);continue
        heading=re.match(r'^(#{1,4})\s+(.*)',line)
        if heading:
            title=re.sub(r'[^\w\s]',' ',heading[2]);title=re.sub(r'\s+',' ',title).strip()
            doc.add_heading(title,level=max(1,len(heading[1])-1));continue
        bullet=re.match(r'^[-*]\s+(.*)',line)
        if bullet:p=doc.add_paragraph(style='List Bullet');inline(p,bullet[1]);continue
        p=doc.add_paragraph();inline(p,line)
        # Keep an introduction with its following diagram, including across blank lines.
        next_line=next((v for v in lines[i:] if v.strip()),'')
        if next_line.startswith('!['):p.paragraph_format.keep_with_next=True


def main():
    OUTDIR.mkdir(exist_ok=True)
    for key,filename,title in ITEMS:
        lines=(ROOT/'docs'/filename).read_text().splitlines()
        body='\n'.join(lines[1:] if lines and lines[0].startswith('# ') else lines)
        doc=Document();sec=doc.sections[0]
        for border in list(doc.styles.element.xpath('.//w:pBdr')):
            border.getparent().remove(border)
        sec.page_width=Inches(8.5);sec.page_height=Inches(11)
        sec.left_margin=sec.right_margin=Inches(.825);sec.top_margin=sec.bottom_margin=Inches(.7)
        for name in ('Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','Heading 4','Caption'):
            doc.styles[name].font.name='Calibri';doc.styles[name].font.color.rgb=RGBColor(0,0,0)
        doc.styles['Normal'].font.size=Pt(11)
        doc.styles['Normal'].paragraph_format.space_after=Pt(7)
        doc.styles['Normal'].paragraph_format.line_spacing=1.1
        doc.styles['Title'].font.size=Pt(25)
        for name,size in [('Heading 1',15),('Heading 2',12.5),('Heading 3',11.5)]:
            st=doc.styles[name];st.font.size=Pt(size);st.font.bold=True;st.paragraph_format.space_before=Pt(13);st.paragraph_format.space_after=Pt(6);st.paragraph_format.keep_with_next=True
        doc.add_paragraph(title,style='Title')
        p=doc.add_paragraph('EVIDENCE cybersecurity investigator | Version 2.3 | Updated 24 September 2026')
        for r in p.runs:r.font.size=Pt(9);r.italic=True
        markdown(doc,body)
        footer=sec.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        r=footer.add_run('EVIDENCE v2.3  |  ');r.font.size=Pt(8)
        field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
        doc.core_properties.title='EVIDENCE '+title;doc.core_properties.author='EVIDENCE project';doc.core_properties.subject=title
        path=OUTDIR/f'EVIDENCE_v2.3_{key}.docx';doc.save(path);print(path.name)

if __name__=='__main__':main()
